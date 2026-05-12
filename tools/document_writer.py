"""Document Writer - Excel/PDF/Word/Markdown."""
from pathlib import Path
from datetime import datetime


class DocumentWriter:
    @staticmethod
    def excel(filename, sheets_data, output_dir="outputs"):
        """sheets_data = {"Sheet1": [["a","b"],["c","d"]], ...}"""
        import openpyxl
        from openpyxl.styles import Font, PatternFill
        Path(output_dir).mkdir(exist_ok=True)
        path = Path(output_dir) / filename
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for name, rows in sheets_data.items():
            ws = wb.create_sheet(name[:31])
            if rows:
                for i, row in enumerate(rows, 1):
                    for j, val in enumerate(row, 1):
                        c = ws.cell(row=i, column=j, value=val)
                        if i == 1:
                            c.font = Font(bold=True, color="FFFFFF")
                            c.fill = PatternFill("solid", fgColor="1F4E78")
                for col in ws.columns:
                    ml = max((len(str(c.value)) for c in col if c.value), default=10)
                    ws.column_dimensions[col[0].column_letter].width = min(ml + 2, 50)
        wb.save(str(path))
        return str(path)

    @staticmethod
    def docx(filename, title, paragraphs, output_dir="outputs"):
        from docx import Document
        from docx.shared import Pt, RGBColor
        Path(output_dir).mkdir(exist_ok=True)
        path = Path(output_dir) / filename
        doc = Document()
        h = doc.add_heading(title, 0)
        for run in h.runs:
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x78)
        doc.add_paragraph(f"Tarih: {datetime.now().strftime('%d.%m.%Y %H:%M')}")
        doc.add_paragraph()
        for p in paragraphs:
            if isinstance(p, dict):
                if p.get("heading"):
                    doc.add_heading(p["heading"], level=p.get("level", 1))
                if p.get("text"):
                    doc.add_paragraph(p["text"])
                if p.get("bullet"):
                    for b in p["bullet"]:
                        doc.add_paragraph(b, style='List Bullet')
            else:
                doc.add_paragraph(str(p))
        doc.save(str(path))
        return str(path)

    @staticmethod
    def markdown(filename, content, output_dir="outputs"):
        Path(output_dir).mkdir(exist_ok=True)
        path = Path(output_dir) / filename
        path.write_text(content, encoding='utf-8')
        return str(path)

    @staticmethod
    def pdf_from_text(filename, title, paragraphs, output_dir="outputs"):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfgen import canvas
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import (SimpleDocTemplate, Paragraph,
                                            Spacer)
            Path(output_dir).mkdir(exist_ok=True)
            path = Path(output_dir) / filename
            doc = SimpleDocTemplate(str(path), pagesize=A4)
            styles = getSampleStyleSheet()
            story = [Paragraph(title, styles['Title']), Spacer(1, 12)]
            for p in paragraphs:
                story.append(Paragraph(str(p), styles['Normal']))
                story.append(Spacer(1, 8))
            doc.build(story)
            return str(path)
        except ImportError:
            md = DocumentWriter.markdown(
                filename.replace('.pdf', '.md'),
                f"# {title}\n\n" + "\n\n".join(str(p) for p in paragraphs),
                output_dir)
            return md
