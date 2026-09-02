"""Document Reader - PDF/Excel/Word/CSV reading."""
from pathlib import Path


class DocumentReader:
    @staticmethod
    def read(file_path):
        p = Path(file_path)
        if not p.exists():
            return {"error": f"Dosya yok: {file_path}"}
        ext = p.suffix.lower()
        try:
            if ext == ".pdf":
                return DocumentReader._pdf(p)
            elif ext in (".xlsx", ".xls"):
                return DocumentReader._excel(p)
            elif ext == ".docx":
                return DocumentReader._docx(p)
            elif ext == ".csv":
                return DocumentReader._csv(p)
            elif ext in (".txt", ".md", ".log", ".json"):
                return {"type": "text",
                        "content": p.read_text(encoding='utf-8', errors='ignore')}
            else:
                return {"error": f"Desteklenmeyen format: {ext}"}
        except Exception as e:
            return {"error": str(e)[:200]}

    @staticmethod
    def _pdf(p):
        from pypdf import PdfReader
        reader = PdfReader(str(p))
        pages = []
        for i, pg in enumerate(reader.pages):
            try:
                pages.append({"page": i + 1, "text": pg.extract_text() or ""})
            except:
                pages.append({"page": i + 1, "text": ""})
        full = "\n\n".join(f"--- Sayfa {x['page']} ---\n{x['text']}" for x in pages)
        return {"type": "pdf", "page_count": len(pages),
                "content": full, "pages": pages}

    @staticmethod
    def _excel(p):
        import openpyxl
        wb = openpyxl.load_workbook(str(p), data_only=True)
        sheets = {}
        for name in wb.sheetnames:
            ws = wb[name]
            data = []
            for row in ws.iter_rows(values_only=True):
                data.append(list(row))
            sheets[name] = {"rows": ws.max_row, "cols": ws.max_column,
                            "data": data[:200]}
        return {"type": "excel", "sheet_count": len(sheets), "sheets": sheets}

    @staticmethod
    def _docx(p):
        from docx import Document
        doc = Document(str(p))
        paras = [pa.text for pa in doc.paragraphs if pa.text.strip()]
        tables = []
        for t in doc.tables:
            rows = []
            for r in t.rows:
                rows.append([c.text for c in r.cells])
            tables.append(rows)
        return {"type": "docx", "paragraph_count": len(paras),
                "content": "\n".join(paras), "tables": tables}

    @staticmethod
    def _csv(p):
        import pandas as pd
        df = pd.read_csv(str(p))
        return {"type": "csv", "rows": len(df), "cols": len(df.columns),
                "columns": list(df.columns),
                "preview": df.head(20).to_dict(orient='records'),
                "summary": df.describe(include='all').to_dict()}

    @staticmethod
    def summary(file_path, max_chars=2000):
        d = DocumentReader.read(file_path)
        if "error" in d:
            return d["error"]
        t = d.get("type")
        if t == "pdf":
            return f"PDF — {d['page_count']} sayfa\n\n{d['content'][:max_chars]}"
        if t == "excel":
            out = [f"Excel — {d['sheet_count']} sayfa"]
            for name, s in d["sheets"].items():
                out.append(f"\n[{name}] {s['rows']}x{s['cols']}")
                for row in s["data"][:5]:
                    out.append(str(row))
            return "\n".join(out)[:max_chars]
        if t == "docx":
            return f"Word — {d['paragraph_count']} paragraf\n\n{d['content'][:max_chars]}"
        if t == "csv":
            return (f"CSV — {d['rows']} satır, {d['cols']} sütun\n"
                    f"Sütunlar: {d['columns']}\nÖrnek: {d['preview'][:3]}")
        if t == "text":
            return d["content"][:max_chars]
        return "Okunamadı"
